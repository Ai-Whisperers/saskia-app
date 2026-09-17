"""tests/test_excel_patch.py — Excel PATCH mode + plantilla endpoint tests.

Per Stream C (prelaunch roadmap 2026-09-17).

Covers:
- PATCH on Productos by name match
- PATCH on Productos by sku match
- PATCH on Productos with unknown name (warning, no create)
- PATCH on Clientes with new phone (creates)
- PATCH on Clientes with existing phone (updates)
- PATCH on Ingredientes updates stock_qty (the "I bought flour" scenario)
- PATCH on Ingredientes updates purchase_price_gs (price update scenario)
- PATCH on Recetas updates yield_qty
- Mixed sheet: bad customer row + good product row (partial success)
- End-to-end: download plantilla, edit one product price, upload, verify update
- Mode selector on UI renders PATCH + FULL
- Endpoints return 422 if mode param missing/garbage
"""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select


def _make_wb(sheets: dict[str, list[list]]) -> Workbook:
    """Build an openpyxl Workbook with the given sheet name → rows map.

    First row of each sheet is treated as the header (string).
    """
    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    return wb


def _sheet_rows(wb: Workbook, name: str) -> list[dict]:
    """Read a sheet into a list of dicts (header → key)."""
    ws: Worksheet | None = wb[name] if name in wb.sheetnames else None
    if ws is None:
        return []
    out = []
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return out
    header = [str(h) if h is not None else "" for h in rows[0]]
    for row in rows[1:]:
        if row is None or all(c is None for c in row):
            continue
        out.append({header[i]: row[i] for i in range(min(len(header), len(row)))})
    return out


def _seed_basic(session_factory):
    """Create 2 products, 1 ingredient, 1 recipe, return their IDs/names."""
    from app.rms.models import Ingredient, Product, Recipe

    with session_factory() as s:
        ing = Ingredient(name="Harina", unit="kg", stock_qty=1.0, min_stock_qty=0.5)
        s.add(ing)
        s.flush()
        ing_id = ing.id

        rec = Recipe(name="Muffin", yield_qty=12.0, yield_unit="und")
        s.add(rec)
        s.flush()
        rec_id = rec.id

        p1 = Product(
            name="Muffin Classic",
            portion_label="1 unidad",
            sale_price_gs=5000,
            sku="MUF-CL-001",
        )
        p2 = Product(
            name="Muffin Choco",
            portion_label="1 unidad",
            sale_price_gs=6000,
        )
        s.add_all([p1, p2])
        s.flush()
        p1_id, p2_id = p1.id, p2.id
    return {
        "ing": ("Harina", ing_id),
        "rec": ("Muffin", rec_id),
        "products": [
            ("Muffin Classic", p1_id, "MUF-CL-001"),
            ("Muffin Choco", p2_id, None),
        ],
    }


# ---------------------------------------------------------------------------
# PATCH mode — Productos
# ---------------------------------------------------------------------------


def test_patch_productos_by_name_updates_price(session_factory):
    """Match by name (case-insensitive) → update sale_price_gs."""
    from app.rms.models import Product
    from app.services.import_xlsx import from_workbook

    seeds = _seed_basic(session_factory)
    p_name, p_id, _ = seeds["products"][0]

    wb = _make_wb(
        {
            "Productos": [
                ["name", "sku", "sale_price_gs", "portion_label", "notes"],
                [p_name.upper(), None, "Gs. 7.500", None, "nueva nota"],
            ]
        }
    )

    s = session_factory()
    try:
        result = from_workbook(s, wb, mode="PATCH")
        assert result.products == 1
        # No warnings expected — match by name worked
        assert result.warnings == []
    finally:
        s.close()

    with session_factory() as s:
        p = s.get(Product, p_id)
        assert p.sale_price_gs == 7500
        assert p.notes == "nueva nota"


def test_patch_productos_by_sku_updates_price(session_factory):
    """Match by sku → update sale_price_gs even when name differs."""
    from app.rms.models import Product
    from app.services.import_xlsx import from_workbook

    seeds = _seed_basic(session_factory)
    _, p_id, p_sku = seeds["products"][0]  # has sku

    wb = _make_wb(
        {
            "Productos": [
                ["name", "sku", "sale_price_gs"],
                [None, p_sku, "Gs. 9.500"],
            ]
        }
    )

    s = session_factory()
    try:
        result = from_workbook(s, wb, mode="PATCH")
        assert result.products == 1
        assert result.warnings == []
    finally:
        s.close()

    with session_factory() as s:
        p = s.get(Product, p_id)
        assert p.sale_price_gs == 9500


def test_patch_productos_unknown_name_warns_no_create(session_factory):
    """Unknown name → warning, no row created (no auto-create in PATCH)."""
    from app.rms.models import Product
    from app.services.import_xlsx import from_workbook

    _seed_basic(session_factory)

    wb = _make_wb(
        {
            "Productos": [
                ["name", "sale_price_gs"],
                ["No Existe Tal Producto", "Gs. 5.000"],
            ]
        }
    )

    s = session_factory()
    try:
        result = from_workbook(s, wb, mode="PATCH")
        assert result.products == 0  # nothing updated
        assert any("No Existe" in w for w in result.warnings)
    finally:
        s.close()

    with session_factory() as s:
        # Product count unchanged
        assert s.scalar(select(Product).order_by(Product.id).limit(1).offset(0)) is not None
        all_products = s.scalars(select(Product)).all()
        # 2 from seed_basic; no new "No Existe" product
        assert len(all_products) == 2
        assert not any(p.name == "No Existe Tal Producto" for p in all_products)


def test_patch_productos_empty_cell_leaves_field_unchanged(session_factory):
    """Empty cell in PATCH → field unchanged (the "edit only what you want" promise)."""
    from app.rms.models import Product
    from app.services.import_xlsx import from_workbook

    seeds = _seed_basic(session_factory)
    p_name, p_id, _ = seeds["products"][0]

    # Row with only `name` populated → no fields changed
    wb = _make_wb(
        {
            "Productos": [
                ["name", "sale_price_gs", "notes"],
                [p_name, None, None],  # all blank except name
            ]
        }
    )

    s = session_factory()
    try:
        result = from_workbook(s, wb, mode="PATCH")
        assert result.products == 1
    finally:
        s.close()

    with session_factory() as s:
        p = s.get(Product, p_id)
        # sale_price unchanged
        assert p.sale_price_gs == 5000


# ---------------------------------------------------------------------------
# PATCH mode — Clientes
# ---------------------------------------------------------------------------


def test_patch_clientes_new_phone_creates_customer(session_factory):
    """New phone in PATCH Clientes → create new Customer with phone + name."""
    from app.rms.models import Customer
    from app.services.import_xlsx import from_workbook

    wb = _make_wb(
        {
            "Clientes": [
                ["phone", "name", "email", "cedula", "notes"],
                ["+595991234567", "Saskia Boer", "saskia@example.com", "1234567", None],
            ]
        }
    )

    s = session_factory()
    try:
        result = from_workbook(s, wb, mode="PATCH")
        assert result.customers == 1
        assert result.warnings == []
    finally:
        s.close()

    with session_factory() as s:
        c = s.scalar(select(Customer).where(Customer.phone == "+595991234567"))
        assert c is not None
        assert c.name == "Saskia Boer"
        assert c.email == "saskia@example.com"
        assert c.cedula == "1234567"


def test_patch_clientes_existing_phone_updates_fields(session_factory):
    """Existing phone → update name, email, etc."""
    from app.rms.models import Customer
    from app.services.import_xlsx import from_workbook

    # Seed
    with session_factory() as s:
        c = Customer(name="Old Name", phone="+595991111111", email="old@example.com")
        s.add(c)
        s.flush()
        c_id = c.id

    wb = _make_wb(
        {
            "Clientes": [
                ["phone", "name", "email"],
                ["+595991111111", "Updated Name", "new@example.com"],
            ]
        }
    )

    s = session_factory()
    try:
        result = from_workbook(s, wb, mode="PATCH")
        assert result.customers == 1
    finally:
        s.close()

    with session_factory() as s:
        c = s.get(Customer, c_id)
        assert c.name == "Updated Name"
        assert c.email == "new@example.com"


def test_patch_clientes_missing_phone_skipped(session_factory):
    """Row without phone → warning, no row touched."""
    from app.services.import_xlsx import from_workbook

    wb = _make_wb(
        {
            "Clientes": [
                ["phone", "name"],
                [None, "Some Guy"],
            ]
        }
    )

    s = session_factory()
    try:
        result = from_workbook(s, wb, mode="PATCH")
        assert result.customers == 0
        assert any("sin phone" in w for w in result.warnings)
    finally:
        s.close()


# ---------------------------------------------------------------------------
# PATCH mode — Ingredientes
# ---------------------------------------------------------------------------


def test_patch_ingredientes_updates_stock_qty(session_factory):
    """'I just bought flour' workflow — operator updates stock_qty via PATCH."""
    from app.rms.models import Ingredient
    from app.services.import_xlsx import from_workbook

    seeds = _seed_basic(session_factory)
    ing_name, ing_id = seeds["ing"]

    wb = _make_wb(
        {
            "Ingredientes": [
                ["name", "stock_qty"],
                [ing_name, 25.5],
            ]
        }
    )

    s = session_factory()
    try:
        result = from_workbook(s, wb, mode="PATCH")
        assert result.ingredients == 1
        assert result.warnings == []
    finally:
        s.close()

    with session_factory() as s:
        ing = s.get(Ingredient, ing_id)
        assert ing.stock_qty == 25.5


def test_patch_ingredientes_price_update_sets_timestamp(session_factory):
    """Purchase price update → automatic purchase_price_updated_at = now."""
    from datetime import datetime, timedelta

    from app.rms.models import Ingredient
    from app.services.import_xlsx import from_workbook

    seeds = _seed_basic(session_factory)
    ing_name, ing_id = seeds["ing"]

    before = datetime.utcnow() - timedelta(minutes=1)

    wb = _make_wb(
        {
            "Ingredientes": [
                ["name", "purchase_price_gs"],
                [ing_name, "Gs. 12.500"],
            ]
        }
    )

    s = session_factory()
    try:
        result = from_workbook(s, wb, mode="PATCH")
        assert result.ingredients == 1
    finally:
        s.close()

    with session_factory() as s:
        ing = s.get(Ingredient, ing_id)
        assert ing.purchase_price_gs == 12500
        assert ing.purchase_price_updated_at is not None
        assert ing.purchase_price_updated_at >= before


def test_patch_ingredientes_unknown_name_warns(session_factory):
    """Unknown ingredient name → warning, no create."""
    from app.services.import_xlsx import from_workbook

    _seed_basic(session_factory)

    wb = _make_wb(
        {
            "Ingredientes": [
                ["name", "stock_qty"],
                ["No Existe Ingrediente", 5.0],
            ]
        }
    )

    s = session_factory()
    try:
        result = from_workbook(s, wb, mode="PATCH")
        assert result.ingredients == 0
        assert any("No Existe" in w for w in result.warnings)
    finally:
        s.close()


# ---------------------------------------------------------------------------
# PATCH mode — Recetas
# ---------------------------------------------------------------------------


def test_patch_recetas_updates_yield_qty(session_factory):
    """PATCH Recetas updates yield_qty by name match."""
    from app.rms.models import Recipe
    from app.services.import_xlsx import from_workbook

    seeds = _seed_basic(session_factory)
    rec_name, rec_id = seeds["rec"]

    wb = _make_wb(
        {
            "Recetas": [
                ["name", "yield_qty", "prep_minutes"],
                [rec_name, 24.0, 90],
            ]
        }
    )

    s = session_factory()
    try:
        result = from_workbook(s, wb, mode="PATCH")
        assert result.recipes == 1
    finally:
        s.close()

    with session_factory() as s:
        rec = s.get(Recipe, rec_id)
        assert rec.yield_qty == 24.0
        assert rec.prep_minutes == 90


# ---------------------------------------------------------------------------
# Mixed-scenario partial-success
# ---------------------------------------------------------------------------


def test_patch_partial_success_bad_customer_with_good_product(session_factory):
    """A bad Clientes row + good Productos row → partial success, bad logged."""
    from app.rms.models import Customer, Product
    from app.services.import_xlsx import from_workbook

    seeds = _seed_basic(session_factory)
    p_name, p_id, _ = seeds["products"][0]

    wb = _make_wb(
        {
            "Productos": [
                ["name", "sale_price_gs"],
                [p_name, "Gs. 8.000"],
            ],
            "Clientes": [
                ["phone", "name"],
                [None, "Sin Telefono"],  # invalid — no phone
                ["+59599999999", "Cliente OK"],  # valid
            ],
        }
    )

    s = session_factory()
    try:
        result = from_workbook(s, wb, mode="PATCH")
        # Product updated, 1 customer created (bad one skipped)
        assert result.products == 1
        assert result.customers == 1
        assert any("sin phone" in w for w in result.warnings)
    finally:
        s.close()

    with session_factory() as s:
        p = s.get(Product, p_id)
        assert p.sale_price_gs == 8000
        c = s.scalar(select(Customer).where(Customer.phone == "+59599999999"))
        assert c is not None
        assert c.name == "Cliente OK"


# ---------------------------------------------------------------------------
# End-to-end: plantilla → edit → upload
# ---------------------------------------------------------------------------


def test_plantilla_prepopulates_products(session_factory):
    """patch_plantilla_bytes() includes all current products with their prices."""
    from app.services.export_xlsx import patch_plantilla_bytes

    seeds = _seed_basic(session_factory)
    with session_factory() as s:
        data = patch_plantilla_bytes(s)

    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "Productos" in wb.sheetnames
    rows = _sheet_rows(wb, "Productos")
    # Both products present
    names = {r["name"] for r in rows}
    assert seeds["products"][0][0] in names
    assert seeds["products"][1][0] in names


def test_plantilla_clientes_header_only(session_factory):
    """Plantilla Clientes sheet has only the header, no body rows."""
    from app.services.export_xlsx import patch_plantilla_bytes

    _seed_basic(session_factory)
    with session_factory() as s:
        data = patch_plantilla_bytes(s)

    wb = openpyxl.load_workbook(io.BytesIO(data))
    rows = _sheet_rows(wb, "Clientes")
    assert rows == []
    # header row exists
    ws = wb["Clientes"]
    header = [c.value for c in ws[1]]
    assert header == ["phone", "name", "email", "cedula", "notes"]


def test_end_to_end_plantilla_edit_upload_updates_price(session_factory):
    """Download plantilla → edit one price → upload via PATCH → confirm update."""
    from app.rms.models import Product
    from app.services.export_xlsx import patch_plantilla_bytes
    from app.services.import_xlsx import from_workbook

    seeds = _seed_basic(session_factory)
    p_name, p_id, _ = seeds["products"][0]

    # Step 1: download (within the same session scope as the seeded data)
    with session_factory() as s:
        data = patch_plantilla_bytes(s)

    wb = openpyxl.load_workbook(io.BytesIO(data))

    # Step 2: edit the price
    ws = wb["Productos"]
    # Locate the row by name (column A) and write to column C
    header = [c.value for c in ws[1]]
    name_col = header.index("name") + 1  # 1-indexed
    price_col = header.index("sale_price_gs") + 1
    edited = False
    for row_idx in range(2, ws.max_row + 1):
        if ws.cell(row=row_idx, column=name_col).value == p_name:
            ws.cell(row=row_idx, column=price_col, value="Gs. 12.345")
            edited = True
            break
    assert edited, f"could not find product {p_name!r} in plantilla Productos sheet"

    # Step 3: upload via PATCH
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb2 = openpyxl.load_workbook(buf, data_only=True, read_only=True)

    s = session_factory()
    try:
        result = from_workbook(s, wb2, mode="PATCH")
        assert result.products == 1
        assert result.warnings == []
    finally:
        s.close()

    # Step 4: verify
    with session_factory() as s:
        p = s.get(Product, p_id)
        assert p.sale_price_gs == 12345


# ---------------------------------------------------------------------------
# Mode validation (422)
# ---------------------------------------------------------------------------


def test_import_endpoint_with_invalid_mode_returns_422(client):
    """POST /excel/importar?mode=GARBAGE → 422."""
    # Build a minimal valid xlsx so we isolate the mode-validation path
    wb = _make_wb({"Productos": [["name", "sale_price_gs"]]})
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = client.post(
        "/excel/importar?mode=GARBAGE",
        files={
            "file": (
                "test.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 422
    assert "mode" in r.text or "GARBAGE" in r.text or "invalid" in r.text.lower()


def test_excel_page_renders_mode_radio_buttons(client):
    """GET /excel → page has PATCH and FULL radio inputs."""
    r = client.get("/excel")
    assert r.status_code == 200
    assert 'value="PATCH"' in r.text
    assert 'value="FULL"' in r.text
    assert "Actualizar por nombre" in r.text or "PATCH" in r.text
    assert "Reemplazar todo" in r.text or "FULL" in r.text


def test_excel_page_renders_plantilla_link(client):
    """GET /excel → page has link to /excel/plantilla."""
    r = client.get("/excel")
    assert r.status_code == 200
    assert "/excel/plantilla" in r.text


def test_plantilla_endpoint_returns_xlsx(client, session_factory):
    """GET /excel/plantilla → 200 with .xlsx body and friendly filename."""
    _seed_basic(session_factory)
    r = client.get("/excel/plantilla")
    assert r.status_code == 200
    assert r.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    # filename pattern: saskia-import-YYYYMMDD.xlsx
    cd = r.headers.get("content-disposition", "")
    assert "saskia-import-" in cd
    assert ".xlsx" in cd
    # Body is a real xlsx file (PK magic bytes)
    assert r.content[:2] == b"PK"
